#!/usr/bin/env python3
"""Apply final institutional pagination/title fixes after DOCX rendering.

The consolidated RF baseline deliberately starts its v1.1 incorporation on a
new page. This prevents the governance close-out from becoming an accidental
near-empty trailing page. The traceability matrix title is also canonicalized
here because its v1.0 source has a longer historical title.
"""
from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

from docx import Document

CANONICAL_MATRIX_TITLE = "Matriz de Rastreabilidade - Jornada do Cidadão - Fase 1"
RF_ADDITIONS_HEADING = "Requisitos funcionais incorporados na v1.1"


def convert_pdf(docx: Path) -> Path:
    pdf = docx.with_suffix(".pdf")
    if pdf.exists():
        pdf.unlink()
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx.parent),
            str(docx),
        ],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    if not pdf.is_file() or pdf.stat().st_size < 1000:
        raise RuntimeError(f"invalid PDF after finalization: {pdf}")
    return pdf


def finalize_rf(req: Path) -> None:
    path = req / "02_Requisitos_Funcionais_Jornada_v1.1.docx"
    doc = Document(path)
    matches = [p for p in doc.paragraphs if p.text.strip() == RF_ADDITIONS_HEADING]
    if len(matches) != 1:
        raise RuntimeError(
            f"expected exactly one RF additions heading, found {len(matches)}"
        )
    matches[0].paragraph_format.page_break_before = True
    doc.save(path)
    convert_pdf(path)


def finalize_matrix(req: Path) -> None:
    md = req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.md"
    lines = md.read_text(encoding="utf-8").splitlines()
    if not lines or not lines[0].startswith("# "):
        raise RuntimeError("traceability matrix Markdown title not found")
    lines[0] = f"# {CANONICAL_MATRIX_TITLE}"
    md.write_text("\n".join(lines) + "\n", encoding="utf-8")

    path = req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.docx"
    doc = Document(path)
    if not doc.paragraphs:
        raise RuntimeError("traceability matrix DOCX has no title paragraph")
    title = doc.paragraphs[0]
    title.text = CANONICAL_MATRIX_TITLE
    title.style = doc.styles["Heading 1"]
    title.paragraph_format.keep_with_next = True
    doc.save(path)
    convert_pdf(path)


def validate(req: Path) -> None:
    rf = Document(req / "02_Requisitos_Funcionais_Jornada_v1.1.docx")
    matches = [p for p in rf.paragraphs if p.text.strip() == RF_ADDITIONS_HEADING]
    if len(matches) != 1 or not matches[0].paragraph_format.page_break_before:
        raise RuntimeError("RF v1.1 incorporation must start on a new page")

    matrix_md = (req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.md").read_text(
        encoding="utf-8"
    )
    if matrix_md.splitlines()[0] != f"# {CANONICAL_MATRIX_TITLE}":
        raise RuntimeError("matrix Markdown title is not canonical")
    matrix = Document(req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.docx")
    if matrix.paragraphs[0].text.strip() != CANONICAL_MATRIX_TITLE:
        raise RuntimeError("matrix DOCX title is not canonical")

    print(
        "REQUIREMENTS LAYOUT FINALIZATION: OK "
        "(RF additions start on new page; matrix title canonical)"
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--root", type=Path, default=Path(__file__).resolve().parents[2]
    )
    args = parser.parse_args()
    req = args.root.resolve() / "Documentos" / "Requisitos"
    finalize_rf(req)
    finalize_matrix(req)
    validate(req)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
