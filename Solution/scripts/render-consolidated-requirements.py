#!/usr/bin/env python3
"""Render the consolidated requirements baseline to institutional DOCX/PDF."""
from __future__ import annotations

import argparse
import importlib.util
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt


def load_helper(root: Path):
    path = root / "Solution" / "scripts" / "generate-document-deliverables.py"
    spec = importlib.util.spec_from_file_location("jornada_document_helper", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("cannot load document helper")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def convert_pdf(docx: Path) -> Path:
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    pdf = docx.with_suffix(".pdf")
    if not pdf.is_file() or pdf.stat().st_size < 1000:
        raise RuntimeError(f"invalid PDF: {pdf}")
    return pdf


def matrix_widths(ncol: int) -> list[float]:
    if ncol == 5:
        return [2.4, 5.3, 5.3, 5.8, 8.1]
    if ncol == 4:
        return [5.0, 4.4, 5.8, 11.7]
    if ncol == 3:
        return [5.0, 7.0, 14.9]
    return [26.9 / ncol] * ncol


def build_matrix(helper, md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.top_margin = Cm(1.25)
    sec.bottom_margin = Cm(1.25)
    sec.left_margin = Cm(1.4)
    sec.right_margin = Cm(1.4)

    for style_name, size, bold in [
        ("Normal", 8.6, False),
        ("Heading 1", 15.0, True),
        ("Heading 2", 11.5, True),
        ("Heading 3", 10.2, True),
        ("Footer", 8.0, False),
    ]:
        style = doc.styles[style_name]
        helper.set_font(style, size=size, bold=bold)
        if style_name == "Normal":
            style.paragraph_format.space_after = Pt(2)
            style.paragraph_format.line_spacing = 1.0
        elif style_name.startswith("Heading"):
            style.paragraph_format.space_before = Pt(5)
            style.paragraph_format.space_after = Pt(2)
    helper.add_page_number(sec, size=8)

    def add_table(rows: list[list[str]]) -> None:
        if not rows:
            return
        ncol = len(rows[0])
        widths = matrix_widths(ncol)
        table = doc.add_table(rows=1, cols=ncol)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        header = table.rows[0]
        helper.set_header_repeat(header)
        font_size = 7.0 if ncol >= 5 else 7.7
        for j, value in enumerate(rows[0]):
            cell = header.cells[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            helper.set_cell_width(cell, widths[j])
            helper.set_cell_margins(cell, top=45, start=55, bottom=45, end=55)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            helper.add_inline(para, value, size=font_size)
            for run in para.runs:
                run.bold = True
        for values in rows[1:]:
            if len(values) != ncol:
                raise RuntimeError(f"matrix row has {len(values)} columns; expected {ncol}: {values}")
            row = table.add_row()
            for j, value in enumerate(values):
                cell = row.cells[j]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                helper.set_cell_width(cell, widths[j])
                helper.set_cell_margins(cell, top=38, start=50, bottom=38, end=50)
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                helper.add_inline(para, value, size=font_size)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)

    idx = 0
    while idx < len(lines):
        line = lines[idx]
        if not line.strip():
            idx += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            helper.add_inline(p, line[2:])
            idx += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            helper.add_inline(p, line[3:])
            idx += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            helper.add_inline(p, line[4:])
            idx += 1
            continue
        if line.startswith("|"):
            block: list[str] = []
            while idx < len(lines) and lines[idx].startswith("|"):
                block.append(lines[idx])
                idx += 1
            rows: list[list[str]] = []
            for pos, raw in enumerate(block):
                values = [v.strip() for v in raw.strip().strip("|").split("|")]
                if pos == 1 and all(re.fullmatch(r":?-{3,}:?", v) for v in values):
                    continue
                rows.append(values)
            add_table(rows)
            continue
        p = doc.add_paragraph()
        helper.add_inline(p, line.rstrip())
        idx += 1

    doc.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[2])
    args = parser.parse_args()
    root = args.root.resolve()
    req = root / "Documentos" / "Requisitos"
    helper = load_helper(root)

    with tempfile.TemporaryDirectory(prefix="jornada-req-render-") as tmp:
        reference = Path(tmp) / "reference_portrait.docx"
        helper.make_reference(reference, landscape=False)
        for stem in [
            "00_Indice_Mestre_Requisitos_Jornada_v1.1",
            "02_Requisitos_Funcionais_Jornada_v1.1",
            "03_Requisitos_Nao_Funcionais_Jornada_v1.1",
        ]:
            md = req / f"{stem}.md"
            docx = req / f"{stem}.docx"
            subprocess.run(
                ["pandoc", str(md), "--reference-doc=" + str(reference), "-o", str(docx)],
                check=True,
            )
            convert_pdf(docx)

    matrix_md = req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.md"
    matrix_docx = matrix_md.with_suffix(".docx")
    build_matrix(helper, matrix_md, matrix_docx)
    convert_pdf(matrix_docx)

    required = {
        "00_Indice_Mestre_Requisitos_Jornada_v1.1.md": ["um único documento por número", "SolutionSchema v3.70"],
        "02_Requisitos_Funcionais_Jornada_v1.1.md": ["RF-001", "RF-050", "RF-051", "RF-056"],
        "03_Requisitos_Nao_Funcionais_Jornada_v1.1.md": ["RNF-001", "RNF-033", "RNF34-A", "RNF34-D"],
        "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.md": ["RF-001", "RF-051", "RNF-001", "RNF34-C"],
    }
    for name, markers in required.items():
        md = req / name
        text = md.read_text(encoding="utf-8")
        for marker in markers:
            if marker not in text:
                raise RuntimeError(f"{name}: missing marker {marker}")
        docx = md.with_suffix(".docx")
        pdf = md.with_suffix(".pdf")
        Document(docx)
        if pdf.stat().st_size < 1000:
            raise RuntimeError(f"invalid PDF: {pdf}")

    matrix_doc = Document(req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.docx")
    if not matrix_doc.tables or any(table.style.name != "Table Grid" for table in matrix_doc.tables):
        raise RuntimeError("traceability matrix must use real Table Grid tables")

    print("REQUIREMENTS RENDER: OK (00/02/03/05 v1.1 DOCX/PDF)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
