#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def run(*args: str, cwd: Path | None = None) -> None:
    subprocess.run(args, cwd=cwd, check=True)


def set_font(style, name: str = "Arial", size: float = 10.5, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def add_page_number(section, size: float = 8.5) -> None:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Página ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    set_font(footer.style, size=size)


def make_reference(path: Path, landscape: bool = False) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Cm(29.7)
        sec.page_height = Cm(21.0)
    else:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)

    styles = doc.styles
    set_font(styles["Normal"], size=10.5)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.06
    for name, size in [
        ("Title", 19),
        ("Subtitle", 11),
        ("Heading 1", 16),
        ("Heading 2", 13),
        ("Heading 3", 11.5),
        ("Heading 4", 10.5),
    ]:
        if name in styles:
            set_font(styles[name], size=size, bold=(name == "Title" or name.startswith("Heading")))
    if "Heading 1" in styles:
        styles["Heading 1"].paragraph_format.space_before = Pt(10)
        styles["Heading 1"].paragraph_format.space_after = Pt(5)
    if "Heading 2" in styles:
        styles["Heading 2"].paragraph_format.space_before = Pt(9)
        styles["Heading 2"].paragraph_format.space_after = Pt(4)
    if "Heading 3" in styles:
        styles["Heading 3"].paragraph_format.space_before = Pt(7)
        styles["Heading 3"].paragraph_format.space_after = Pt(3)
    if "Compact" in styles:
        set_font(styles["Compact"], size=10)
    if "Source Code" in styles:
        set_font(styles["Source Code"], name="Consolas", size=9)
    add_page_number(sec)
    doc.save(path)


def add_inline(paragraph, text: str, size: float | None = None) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run_ = paragraph.add_run(part[2:-2])
            run_.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run_ = paragraph.add_run(part[1:-1])
            run_.font.name = "Consolas"
            run_._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
        else:
            run_ = paragraph.add_run(part)
        if not run_.font.name:
            run_.font.name = "Arial"
        run_._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), run_.font.name)
        if size is not None:
            run_.font.size = Pt(size)


def set_cell_width(cell, cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(cm).emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def set_header_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_margins(cell, top: int = 55, start: int = 70, bottom: int = 55, end: int = 70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def build_matrix(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.top_margin = Cm(1.35)
    sec.bottom_margin = Cm(1.35)
    sec.left_margin = Cm(1.4)
    sec.right_margin = Cm(1.4)

    for style_name, size, bold in [
        ("Normal", 9.3, False),
        ("Heading 1", 16, True),
        ("Heading 2", 12.5, True),
        ("Heading 3", 11, True),
        ("Footer", 8, False),
    ]:
        style = doc.styles[style_name]
        set_font(style, size=size, bold=bold)
        if style_name == "Normal":
            style.paragraph_format.space_after = Pt(3)
            style.paragraph_format.line_spacing = 1.0
        elif style_name.startswith("Heading"):
            style.paragraph_format.space_before = Pt(6)
            style.paragraph_format.space_after = Pt(3)
    add_page_number(sec, size=8)

    def add_table(rows: list[list[str]]) -> None:
        ncol = len(rows[0])
        table = doc.add_table(rows=1, cols=ncol)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        if ncol == 5:
            widths = [2.4, 5.3, 5.3, 5.8, 8.1]
        elif ncol == 4:
            widths = [5.0, 4.4, 5.8, 11.7]
        elif ncol == 3:
            widths = [5.0, 7.0, 14.9]
        else:
            widths = [26.9 / ncol] * ncol
        header = table.rows[0]
        set_header_repeat(header)
        for j, value in enumerate(rows[0]):
            cell = header.cells[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[j])
            set_cell_margins(cell)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            add_inline(para, value, size=8.3)
            for run_ in para.runs:
                run_.bold = True
        for values in rows[1:]:
            row = table.add_row()
            for j, value in enumerate(values):
                cell = row.cells[j]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                set_cell_width(cell, widths[j])
                set_cell_margins(cell)
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                add_inline(para, value, size=8.1)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    idx = 0
    while idx < len(lines):
        line = lines[idx]
        if not line.strip():
            idx += 1
            continue
        if line.startswith("# "):
            para = doc.add_paragraph(style="Heading 1")
            add_inline(para, line[2:])
            idx += 1
            continue
        if line.startswith("## "):
            para = doc.add_paragraph(style="Heading 2")
            add_inline(para, line[3:])
            idx += 1
            continue
        if line.startswith("|"):
            block: list[str] = []
            while idx < len(lines) and lines[idx].startswith("|"):
                block.append(lines[idx])
                idx += 1
            rows: list[list[str]] = []
            for pos, raw in enumerate(block):
                values = [v.strip() for v in raw.strip().strip("|").split("|")]
                if pos == 1 and all(re.fullmatch(r":?-{3,}:?", v) for v in values):
                    continue
                rows.append(values)
            add_table(rows)
            continue
        para = doc.add_paragraph()
        add_inline(para, line.rstrip())
        idx += 1

    doc.save(out_path)


def make_diagrams(out_dir: Path) -> tuple[Path, Path]:
    import matplotlib.pyplot as plt
    from matplotlib.patches import Circle, FancyArrowPatch, Polygon, Rectangle

    out_dir.mkdir(parents=True, exist_ok=True)

    def class_box(ax, x, y, w, h, title, attrs):
        ax.add_patch(Rectangle((x, y), w, h, fill=False, linewidth=1.2))
        header_y = y + h * 0.72
        ax.plot([x, x + w], [header_y, header_y], linewidth=1.0, color="black")
        ax.text(x + w / 2, y + h * 0.86, title, ha="center", va="center", fontsize=12, fontweight="bold")
        ax.text(x + 0.12, y + h * 0.65, "\n".join(attrs), ha="left", va="top", fontsize=9)

    def arrow(ax, p1, p2, text=None, style="-", both=False):
        ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="<->" if both else "->", mutation_scale=12,
                                     linewidth=1.0, linestyle=style, color="black"))
        if text:
            mx = (p1[0] + p2[0]) / 2
            my = (p1[1] + p2[1]) / 2
            ax.text(mx, my + 0.12, text, ha="center", va="bottom", fontsize=8)

    classes = out_dir / "Jornada_Identidade_Linkage_Classes.png"
    fig, ax = plt.subplots(figsize=(12, 7))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 7)
    ax.axis("off")
    class_box(ax, 0.4, 5.1, 2.0, 1.35, "PessoaOrigem", ["pessoa_origem_id", "sistema_origem_id", "codigo_pessoa_origem", "cpf?"])
    class_box(ax, 3.0, 5.0, 2.4, 1.55, "PessoaOrigemProgressiva", ["pessoa_origem_id", "initial_uuid", "canonical_uuid?", "estado"])
    class_box(ax, 6.1, 5.25, 1.55, 1.05, "Pessoa", ["pessoa_uuid"])
    class_box(ax, 9.0, 5.2, 1.7, 1.15, "CpfAncora", ["cpf", "pessoa_uuid"])
    arrow(ax, (2.4, 5.8), (3.0, 5.8), "1 : 1")
    arrow(ax, (5.4, 5.8), (6.1, 5.8), "* : 1  initial_uuid")
    arrow(ax, (5.4, 5.45), (6.1, 5.45), "* : 0..1  canonical_uuid")
    arrow(ax, (9.0, 5.75), (7.65, 5.75), "0..1 : 1  permanente")
    class_box(ax, 0.5, 2.8, 2.0, 1.35, "LinkageRuleSet", ["ruleset_id", "versao", "fingerprint_sha256", "status"])
    class_box(ax, 3.0, 2.95, 1.8, 1.05, "RuleSetPasse", ["ruleset_id", "passe_ordem"])
    class_box(ax, 5.25, 2.7, 2.15, 1.45, "RuleSetPasseCampo", ["ruleset_id", "passe_ordem", "campo_ordem", "atributo"])
    class_box(ax, 8.15, 2.75, 2.0, 1.35, "BlockingChave", ["pessoa_uuid", "atributo", "valor_normalizado", "normalizacao_versao"])
    arrow(ax, (2.5, 3.45), (3.0, 3.45), "1 : 1..*")
    arrow(ax, (4.8, 3.45), (5.25, 3.45), "1 : 1..*")
    arrow(ax, (7.4, 3.45), (8.15, 3.45), "gera candidatos")
    arrow(ax, (9.15, 4.1), (6.95, 5.25), "0..* : 1")
    class_box(ax, 0.4, 0.55, 2.0, 1.25, "ComposicaoPlano", ["decision_id", "policy_version", "evidence_fingerprint"])
    class_box(ax, 3.0, 0.55, 2.0, 1.25, "ComposicaoAplicacao", ["decision_id", "status", "aplicado_em"])
    class_box(ax, 5.65, 0.55, 2.0, 1.25, "ComposicaoPublicacao", ["decision_id", "status", "publicado_em"])
    class_box(ax, 8.3, 0.55, 1.8, 1.25, "GoldPessoa", ["pessoa_uuid?", "estado_identidade", "dados_factuais"])
    class_box(ax, 10.55, 0.7, 1.1, 0.95, "Serving", ["registro_id", "pessoa_uuid?"])
    arrow(ax, (2.4, 1.18), (3.0, 1.18), "1 : 0..1")
    arrow(ax, (5.0, 1.18), (5.65, 1.18), "1 : 0..1")
    arrow(ax, (7.65, 1.18), (8.3, 1.18), "publicacao atomica")
    arrow(ax, (10.1, 1.18), (10.55, 1.18), "projecao")
    ax.plot([5.0, 5.05], [1.8, 4.72], linestyle="--", linewidth=1.0, color="black")
    arrow(ax, (5.05, 4.72), (5.05, 5.0), "altera referencia", style="--")
    ax.text(9.85, 6.65, "Regra: CPF -> UUID é estável; Linkage probabilístico não transfere âncora.", ha="center", va="center", fontsize=9)
    ax.text(1.5, 2.45, "Calibrador publica ruleset imutável;\nAvaliador consome a mesma versão.", ha="center", va="top", fontsize=9)
    ax.text(9.2, 0.18, "Fato e identidade permanecem conceitos distintos.", ha="center", va="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(classes, dpi=180, bbox_inches="tight")
    plt.close(fig)

    activity = out_dir / "Jornada_Resolucao_Identidade_Atividade.png"
    fig, ax = plt.subplots(figsize=(16, 10))
    ax.set_xlim(0, 16)
    ax.set_ylim(-0.45, 10)
    ax.axis("off")

    def box(cx, cy, w, h, text, fs=10):
        ax.add_patch(Rectangle((cx - w / 2, cy - h / 2), w, h, fill=False, linewidth=1.2))
        ax.text(cx, cy, text, ha="center", va="center", fontsize=fs, wrap=True)

    def dia(cx, cy, w, h, text, fs=9.5):
        points = [(cx, cy + h / 2), (cx + w / 2, cy), (cx, cy - h / 2), (cx - w / 2, cy)]
        ax.add_patch(Polygon(points, closed=True, fill=False, linewidth=1.2))
        ax.text(cx, cy, text, ha="center", va="center", fontsize=fs, wrap=True)

    def arr(x1, y1, x2, y2, label=None):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="->", mutation_scale=12,
                                     linewidth=1.1, color="black"))
        if label:
            ax.text((x1 + x2) / 2, (y1 + y2) / 2 + 0.10, label, fontsize=9, ha="center", va="center")

    def poly(points, label=None, labelpos=None):
        for i in range(len(points) - 2):
            (x1, y1), (x2, y2) = points[i], points[i + 1]
            ax.plot([x1, x2], [y1, y2], linewidth=1.1, color="black")
        (x1, y1), (x2, y2) = points[-2], points[-1]
        arr(x1, y1, x2, y2)
        if label and labelpos:
            ax.text(*labelpos, label, fontsize=9, ha="center", va="center")

    ax.add_patch(Circle((8, 9.75), 0.09, fill=True, color="black"))
    box(8, 9.25, 4.6, 0.48, "Receber observação preservada da origem", 10.5)
    box(8, 8.55, 4.6, 0.48, "Validar contrato, proveniência e qualidade", 10.5)
    box(8, 7.85, 4.6, 0.48, "Assegurar initial_uuid idempotente", 10.5)
    dia(8, 7.05, 3.6, 0.78, "CPF válido, confiável\ne não conflitado?", 10)
    arr(8, 9.66, 8, 9.49)
    arr(8, 9.01, 8, 8.79)
    arr(8, 8.31, 8, 8.09)
    arr(8, 7.61, 8, 7.44)
    ax.text(3.2, 6.8, "Caminho determinístico por CPF", fontsize=10.5, fontweight="bold", ha="center")
    dia(3.2, 6.15, 2.8, 0.70, "Existe cpf_ancora?", 9.5)
    box(2.0, 5.25, 2.8, 0.62, "Recuperar UUID\npermanente da âncora", 9.5)
    box(4.6, 5.25, 2.8, 0.62, "Reservar CPF -> UUID\ntransacional e append-only", 9.2)
    box(3.2, 4.25, 3.1, 0.62, "Produzir decisão determinística", 9.7)
    arr(6.2, 7.03, 4.55, 6.45, "sim")
    arr(2.55, 5.90, 2.0, 5.56, "sim")
    arr(3.85, 5.90, 4.6, 5.56, "não")
    arr(2.0, 4.94, 2.75, 4.56)
    arr(4.6, 4.94, 3.65, 4.56)
    ax.text(12.5, 6.8, "Caminho de Linkage versionado", fontsize=10.5, fontweight="bold", ha="center")
    box(12.5, 6.15, 3.4, 0.62, "Carregar modelo e ruleset\nversionados e imutáveis", 9.4)
    box(12.5, 5.25, 3.4, 0.62, "Gerar candidatos por blocking_chave\ncom todos os passes aplicáveis", 9.2)
    dia(12.5, 4.25, 3.5, 0.76, "Universo completo\ne execução íntegra?", 9.4)
    box(14.5, 3.25, 2.7, 0.65, "Falha operacional:\nnão interpretar como ‘sem candidato’\ne não publicar referência", 8.7)
    box(10.5, 3.45, 2.9, 0.62, "Calcular evidências e score\nsegundo modelo homologado", 9.2)
    box(10.5, 2.65, 3.6, 0.75, "Classificar resultado:\núnico seguro -> ASSOCIAÇÃO_EXISTENTE\nnenhum elegível -> NOVA_IDENTIDADE\nambiguidade -> INDEFINIDA", 8.7)
    arr(9.8, 7.03, 11.05, 6.45, "não")
    arr(12.5, 5.84, 12.5, 5.56)
    arr(12.5, 4.94, 12.5, 4.63)
    arr(13.85, 4.0, 14.5, 3.58, "não")
    arr(11.15, 4.0, 10.5, 3.76, "sim")
    arr(10.5, 3.14, 10.5, 3.02)
    ax.add_patch(Circle((15.55, 2.85), 0.12, fill=False, linewidth=1.1))
    ax.add_patch(Circle((15.55, 2.85), 0.065, fill=True, color="black"))
    arr(15.15, 3.05, 15.45, 2.90)
    box(8, 1.85, 4.8, 0.62, "Persistir decision_id + política/ruleset/modelo + evidência, de forma idempotente", 9.2)
    poly([(3.2, 3.94), (3.2, 3.0), (6.2, 3.0), (6.2, 2.16)])
    poly([(10.5, 2.275), (10.5, 2.16), (9.8, 2.16)])
    dia(8, 1.08, 3.0, 0.62, "Decisão autoriza referência?", 9.2)
    arr(8, 1.54, 8, 1.39)
    box(4.6, 0.38, 3.5, 0.55, "Aplicar composição + publicar referência,\nhistórico e Gold/Serving atomicamente", 8.7)
    box(11.4, 0.38, 3.5, 0.55, "Preservar PROVISÓRIA/INDEFINIDA\ne fatos válidos", 8.8)
    poly([(6.5, 1.08), (5.8, 1.08), (5.8, 0.66)], "sim", (6.0, 1.22))
    poly([(9.5, 1.08), (10.2, 1.08), (10.2, 0.66)], "não", (10.0, 1.22))
    box(8, -0.10, 5.8, 0.34, "Registrar auditoria e recibo de execução", 8.5)
    poly([(4.6, 0.105), (4.6, -0.10), (5.1, -0.10)])
    poly([(11.4, 0.105), (11.4, -0.10), (10.9, -0.10)])
    fig.tight_layout(pad=0.35)
    fig.savefig(activity, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return classes, activity


def make_anexo_source(source: Path, target: Path) -> None:
    text = source.read_text(encoding="utf-8")
    marker = "## 2. Definição canônica do schema"
    insert = """### 1.1 Diagramas UML normativos incorporados

![Figura 1 - Diagrama de classes UML da estrutura de Identidade/Linkage](uml/Jornada_Identidade_Linkage_Classes.png){width=9.2in}

![Figura 2 - Diagrama de atividade UML da resolução de identidade](uml/Jornada_Resolucao_Identidade_Atividade.png){width=9.0in}

"""
    if marker not in text:
        raise RuntimeError("Marcador da seção 2 não encontrado no Anexo v1.40")
    if "### 1.1 Diagramas UML normativos incorporados" not in text:
        text = text.replace(marker, insert + marker, 1)
    target.write_text(text, encoding="utf-8")


def pdf_pages(path: Path) -> int:
    result = subprocess.run(["pdfinfo", str(path)], check=True, text=True, stdout=subprocess.PIPE)
    match = re.search(r"^Pages:\s+(\d+)\s*$", result.stdout, re.MULTILINE)
    if not match:
        raise RuntimeError(f"Não foi possível obter paginação de {path}")
    return int(match.group(1))


def validate(root: Path, work: Path) -> None:
    req = root / "Documentos" / "Requisitos"
    docdir = root / "Documentos"
    docx_files = [
        req / "02_Requisitos_Funcionais_Jornada_v1.1.docx",
        req / "03_Requisitos_Nao_Funcionais_Jornada_v1.1.docx",
        req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.docx",
        docdir / "Anexo_Modelo_Fisico_Jornada_v1.40.docx",
    ]
    pdf_files = [p.with_suffix(".pdf") for p in docx_files]
    # Paginação exata não é um invariante semântico: ela varia com o motor de
    # renderização e já divergia dos PDFs versionados (9/6/7/8 versus 2/2/2/8
    # codificados anteriormente). Os intervalos abaixo detectam colapso ou
    # explosão de layout sem rejeitar variação legítima do LibreOffice.
    expected_page_ranges = [(8, 14), (5, 10), (6, 12), (7, 10)]
    for path in docx_files + pdf_files:
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError(f"Artefato ausente/vazio: {path}")
    for path, (minimum, maximum) in zip(pdf_files, expected_page_ranges):
        pages = pdf_pages(path)
        if pages < minimum or pages > maximum:
            raise RuntimeError(
                f"Paginação fora da faixa em {path.name}: {pages}; esperado entre {minimum} e {maximum}"
            )

    matrix = Document(docx_files[2])
    matrix_source = req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.md"
    source_lines = matrix_source.read_text(encoding="utf-8").splitlines()
    expected_table_count = sum(
        1
        for index, line in enumerate(source_lines)
        if line.startswith("|") and (index == 0 or not source_lines[index - 1].startswith("|"))
    )
    if len(matrix.tables) != expected_table_count:
        raise RuntimeError(
            f"Matriz contém {len(matrix.tables)} tabelas; esperado {expected_table_count} conforme a fonte Markdown"
        )
    if any(t.style is None or t.style.name != "Table Grid" for t in matrix.tables):
        raise RuntimeError("Matriz contém tabela sem a grade institucional Table Grid")
    anexo = Document(docx_files[3])
    if len(anexo.inline_shapes) != 2:
        raise RuntimeError("Anexo não contém exatamente as duas figuras UML incorporadas")

    expected_markers = ["RF-056", "RNF34-C", "RF-051", "69 tabelas"]
    for path, marker in zip(docx_files, expected_markers):
        txt = work / (path.stem + ".txt")
        with txt.open("w", encoding="utf-8") as handle:
            subprocess.run(["pandoc", str(path), "-t", "plain"], check=True, text=True, stdout=handle)
        content = txt.read_text(encoding="utf-8")
        if marker not in content:
            raise RuntimeError(f"Marcador obrigatório ausente em {path.name}: {marker}")
        if re.search(r"\b(?:SGM|SPE)\b", content, re.IGNORECASE):
            raise RuntimeError(f"SGM/SPE encontrado em {path.name}")
        if "{width=" in content:
            raise RuntimeError(f"Atributo de geração vazou para {path.name}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Gera DOCX/PDF institucionais vigentes da Jornada")
    parser.add_argument("--root", default=".", help="Raiz do repositório")
    args = parser.parse_args()
    root = Path(args.root).resolve()
    req = root / "Documentos" / "Requisitos"
    docdir = root / "Documentos"
    work = root / ".generated-doc-deliverables"
    if work.exists():
        shutil.rmtree(work)
    (work / "uml").mkdir(parents=True)

    portrait = work / "reference_portrait.docx"
    landscape = work / "reference_landscape.docx"
    make_reference(portrait, landscape=False)
    make_reference(landscape, landscape=True)
    make_diagrams(work / "uml")

    run("pandoc", str(req / "02_Requisitos_Funcionais_Jornada_v1.1.md"), "--reference-doc=" + str(portrait),
        "-o", str(req / "02_Requisitos_Funcionais_Jornada_v1.1.docx"))
    run("pandoc", str(req / "03_Requisitos_Nao_Funcionais_Jornada_v1.1.md"), "--reference-doc=" + str(portrait),
        "-o", str(req / "03_Requisitos_Nao_Funcionais_Jornada_v1.1.docx"))
    build_matrix(req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.md",
                 req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.docx")

    anexo_md = work / "Anexo_Modelo_Fisico_Jornada_v1.40.md"
    make_anexo_source(docdir / "Anexo_Modelo_Fisico_Jornada_v1.40.md", anexo_md)
    run("pandoc", anexo_md.name, "--reference-doc=" + str(landscape),
        "-o", str(docdir / "Anexo_Modelo_Fisico_Jornada_v1.40.docx"), cwd=work)

    docx_files = [
        req / "02_Requisitos_Funcionais_Jornada_v1.1.docx",
        req / "03_Requisitos_Nao_Funcionais_Jornada_v1.1.docx",
        req / "05_Matriz_Rastreabilidade_Requisitos_Jornada_v1.1.docx",
        docdir / "Anexo_Modelo_Fisico_Jornada_v1.40.docx",
    ]
    for path in docx_files:
        run("libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(path.parent), str(path))

    validate(root, work)
    shutil.rmtree(work)
    print("Artefatos DOCX/PDF gerados e validados: paginação dentro das faixas esperadas, matriz tabular e UML incorporada.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
